from dataclasses import dataclass
from hashlib import sha256
from io import BytesIO
from pathlib import PurePath
import re

from openpyxl import load_workbook


POSITION_ALIASES = {
    "后勤": "logistics",
    "保安": "security",
    "保健": "health",
    "保育": "care",
    "厨房": "kitchen",
    "食堂": "kitchen",
    "业务部门": "business",
    "业务人员": "business",
}
FREQUENCIES = {
    "每日": "daily", "每天": "daily", "日": "daily", "每天": "daily",
    "每周": "weekly", "每周一": "weekly", "每周二": "weekly", "每周三": "weekly",
    "每周四": "weekly", "每周五": "weekly", "每月": "monthly", "月": "monthly",
    "学期": "semester", "每学期": "semester", "学年": "semester",
}
HEADER_KEYS = {
    "frequency": ("周期", "频率"),
    "time_range": ("时间段", "时间", "时段"),
    "task": ("工作职责", "岗位职责", "职责", "工作内容", "工作事项", "任务"),
    "detail": ("具体工作内容", "工作要求", "具体要求", "工作细则", "要求"),
}


@dataclass(frozen=True)
class ParsedRow:
    position_code: str | None
    frequency: str | None
    time_range: str | None
    task_label: str | None
    content: str
    source_sheet: str
    source_row: int
    needs_review: bool


@dataclass(frozen=True)
class ParsedWorkbook:
    file_hash: str
    rows: list[ParsedRow]


def _resolve_position(filename: str, sheet_name: str) -> str | None:
    for source in (sheet_name, PurePath(filename).stem):
        matches = {code for label, code in POSITION_ALIASES.items() if label in source}
        if matches:
            return next(iter(matches)) if len(matches) == 1 else None
    return None


def _normal(value) -> str:
    return " ".join(str(value).replace("\n", " ").split()) if value is not None else ""


def _cell(cells: list[str], indexes: dict[str, int], key: str) -> str:
    index = indexes.get(key)
    return cells[index] if index is not None and 0 <= index < len(cells) else ""


def _frequency(value: str) -> str | None:
    normalized = value.strip()
    if normalized in FREQUENCIES:
        return FREQUENCIES[normalized]
    weekday_match = re.fullmatch(r"每周([一二三四五])", normalized)
    return "weekly" if weekday_match else None


def _find_header(rows: list[tuple], max_scan: int = 12) -> tuple[int | None, dict[str, int]]:
    for row_index, row in enumerate(rows[:max_scan]):
        indexes = {}
        for index, value in enumerate(row):
            label = _normal(value)
            for key, aliases in HEADER_KEYS.items():
                if label in aliases:
                    indexes[key] = index
        if "task" in indexes and ("detail" in indexes or "frequency" in indexes):
            return row_index, indexes
    return None, {}


def parse_workbook(file_bytes: bytes, filename: str) -> ParsedWorkbook:
    if PurePath(filename).suffix.lower() != ".xlsx":
        raise ValueError("only .xlsx workbooks are supported")
    try:
        workbook = load_workbook(BytesIO(file_bytes), data_only=True, read_only=True)
    except Exception as exc:
        raise ValueError("invalid .xlsx workbook") from exc
    parsed = []
    for sheet in workbook.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        header_index, indexes = _find_header(rows)
        position = _resolve_position(filename, sheet.title)
        data_start = header_index + 1 if header_index is not None else 0
        for row_index in range(data_start, len(rows)):
            values = rows[row_index]
            cells = [_normal(value) for value in values]
            if not any(cells):
                continue
            frequency_raw = _cell(cells, indexes, "frequency")
            time_range = _cell(cells, indexes, "time_range")
            task_label = _cell(cells, indexes, "task")
            detail = _cell(cells, indexes, "detail")
            content = "\n".join(value for value in [task_label, detail] if value)
            if not content:
                content = " | ".join(cells)
            frequency = _frequency(frequency_raw)
            needs_review = not (position and frequency and task_label and content)
            parsed.append(ParsedRow(position, frequency, time_range or None, task_label or None, content, sheet.title, row_index + 1, needs_review))
    return ParsedWorkbook(sha256(file_bytes).hexdigest(), parsed)
